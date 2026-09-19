import os
import io
import logging
from typing import Tuple, Optional

logger = logging.getLogger("agent-service.extractor")

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    from PIL import Image
    import pytesseract
except ImportError:
    Image = None
    pytesseract = None

try:
    import docx
except ImportError:
    docx = None


class DocumentTextExtractor:
    """
    Multi-format document text extractor.
    Extracts text directly where possible (PDF, DOCX, TXT),
    runs OCR (Tesseract) on images and scanned documents,
    and returns extraction method and status without fabricating missing content.
    """

    @classmethod
    def extract_from_bytes(cls, file_bytes: bytes, filename: str, mime_type: str = "") -> Tuple[str, str, str]:
        """
        Returns: (extracted_text, method, status)
        method: 'DIRECT_PARSE' | 'OCR' | 'FALLBACK' | 'NONE'
        status: 'READY' | 'NEEDS_REVIEW' | 'FAILED'
        """
        if not file_bytes:
            return "Not detected", "NONE", "FAILED"

        ext = filename.lower().split(".")[-1] if "." in filename else ""

        # 1. Plain text and CSV
        if ext in ["txt", "csv"] or "text/plain" in mime_type or "text/csv" in mime_type:
            try:
                text = file_bytes.decode("utf-8", errors="replace").strip()
                if text:
                    return text, "DIRECT_PARSE", "READY"
                return "Not detected", "NONE", "FAILED"
            except Exception as e:
                logger.warning(f"Error decoding text file: {e}")
                return "Not detected", "NONE", "FAILED"

        # 2. PDF documents
        if ext == "pdf" or "application/pdf" in mime_type:
            return cls._extract_pdf(file_bytes)

        # 3. Microsoft Word DOCX
        if ext == "docx" or "wordprocessingml" in mime_type:
            return cls._extract_docx(file_bytes)

        # 4. Images (JPG, JPEG, PNG) -> OCR
        if ext in ["jpg", "jpeg", "png"] or "image/" in mime_type:
            return cls._extract_image_ocr(file_bytes)

        # 5. Fallback string decode attempt
        try:
            sample = file_bytes[:4096].decode("utf-8", errors="ignore").strip()
            if len(sample) > 50:
                return sample, "FALLBACK", "NEEDS_REVIEW"
        except Exception:
            pass

        return "Not detected", "NONE", "FAILED"

    @classmethod
    def _extract_pdf(cls, file_bytes: bytes) -> Tuple[str, str, str]:
        if pypdf is None:
            return "Not detected", "NONE", "FAILED"
        try:
            stream = io.BytesIO(file_bytes)
            reader = pypdf.PdfReader(stream)
            pages_text = []
            for i, page in enumerate(reader.pages):
                txt = page.extract_text() or ""
                if txt.strip():
                    pages_text.append(txt.strip())

            full_text = "\n\n".join(pages_text).strip()

            # If text is substantial, direct parse succeeded
            if len(full_text) >= 25:
                return full_text, "DIRECT_PARSE", "READY"

            # Scanned PDF: attempt OCR on embedded images if Tesseract is available
            if pytesseract and Image:
                ocr_texts = []
                for page in reader.pages:
                    for img_obj in page.images:
                        try:
                            pil_img = Image.open(io.BytesIO(img_obj.data))
                            ocr_txt = pytesseract.image_to_string(pil_img).strip()
                            if ocr_txt:
                                ocr_texts.append(ocr_txt)
                        except Exception as ocr_err:
                            logger.debug(f"OCR image extraction error: {ocr_err}")
                if ocr_texts:
                    return "\n\n".join(ocr_texts).strip(), "OCR", "READY"

            return full_text if full_text else "Not detected", "DIRECT_PARSE" if full_text else "NONE", "NEEDS_REVIEW" if full_text else "FAILED"
        except Exception as e:
            logger.warning(f"Failed to extract text from PDF: {e}")
            return "Not detected", "NONE", "FAILED"

    @classmethod
    def _extract_docx(cls, file_bytes: bytes) -> Tuple[str, str, str]:
        if docx is None:
            return "Not detected", "NONE", "FAILED"
        try:
            stream = io.BytesIO(file_bytes)
            doc = docx.Document(stream)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            # Also extract table text
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if row_cells:
                        table_texts.append(" | ".join(row_cells))
            full = "\n".join(paragraphs + table_texts).strip()
            if full:
                return full, "DIRECT_PARSE", "READY"
            return "Not detected", "NONE", "FAILED"
        except Exception as e:
            logger.warning(f"Failed to extract text from DOCX: {e}")
            return "Not detected", "NONE", "FAILED"

    @classmethod
    def _extract_image_ocr(cls, file_bytes: bytes) -> Tuple[str, str, str]:
        if Image is None or pytesseract is None:
            return "Not detected", "NONE", "FAILED"
        try:
            stream = io.BytesIO(file_bytes)
            img = Image.open(stream)
            # Real OCR execution
            text = pytesseract.image_to_string(img).strip()
            if text:
                return text, "OCR", "READY"
            return "Not detected", "OCR", "NEEDS_REVIEW"
        except Exception as e:
            logger.warning(f"Image OCR failed: {e}")
            return "Not detected", "OCR", "FAILED"

